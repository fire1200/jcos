#!/usr/bin/env python3
"""4주 실행 캘린더(calendar/4week_calendar.xlsx)와 KPI 추적 시트(tools/kpi_tracker.xlsx)를 생성한다.

실행: python3 scripts/build_workbooks.py
근거 문서: plans/youtube_baby_sleep_exposure_v2_20260916.md 4장·5장
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.workbook.properties import CalcProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1B2A49")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFFFF0")
INPUT_FONT = Font(name=FONT, color="0000FF")
BODY_FONT = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

START = dt.date(2026, 9, 21)  # 월요일. 첫 주 시작일, 필요 시 시트에서 수정


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------- 캘린더
WEEK_THEMES = {
    1: "채널 개설·이관",
    2: "신규 콘텐츠 + A/B",
    3: "글로벌·첫 7일 루틴",
    4: "리뷰·재작성",
}
# (weekday 0=월, 유형, 작업, 산출물/측정)
PLAN = {
    1: [
        (0, "개설", "브랜드 계정으로 새 채널 개설, 채널명·핸들·설명·키워드·시청자층(아동용 아님)·업로드 기본값 설정", "channel_split_plan 1장 체크"),
        (0, "정리", "@PPAIEO 수면음악 39편 일부공개 전환, 재생목록 삭제, 채널 키워드 정리", "일부공개 39편"),
        (1, "세팅", "새 채널 재생목록 5개 생성, 배너·프로필 업로드", "재생목록 링크"),
        (1, "업로드", "롱폼 1번(엄마 허밍 3시간) 재작업 후 예약 공개 20:00, 썸네일 A/B 3종", "CTR 48h"),
        (2, "제작", "롱폼 2번(무가사 3시간) 제목·썸네일·다국어 메타 재작업", "check_metadata.py 통과"),
        (3, "제작", "30분 무가사 첫 10분 재편집(오디오 사양 2장)", "재편집 파일"),
        (4, "업로드", "롱폼 2번 예약 공개 20:00, 최종화면 1번↔2번 상호 연결", "CTR 48h"),
        (5, "유통", "2분짜리 33곡을 허밍·무가사 앨범 2장으로 정리, 유통사·팟캐스트 호스팅 계정 개설", "트랙 목록, 계정"),
        (6, "측정", "새 채널 첫 주 지표·@PPAIEO 트래픽 소스 기준선 기록", "KPI 시트 W0"),
    ],
    2: [
        (0, "커뮤니티", "설문 결과 발표 + 화요일 공개 예고", "게시글"),
        (1, "업로드", "롱폼 #1 공개 (새 제목 공식, 썸네일 A/B 3종)", "CTR 48h"),
        (2, "쇼츠", "쇼츠 2편, 관련 동영상 링크 → 롱폼 #1", "쇼츠→롱폼 클릭"),
        (3, "커뮤니티", "팁 카드 1건 + 댓글 전체 답변", "답변율 100%"),
        (4, "업로드", "롱폼 #2 공개 (다른 시리즈, 썸네일 A/B)", "CTR 48h"),
        (5, "쇼츠", "쇼츠 3편 (롱폼 #2 클립 + 훅형)", "쇼츠 조회수"),
        (6, "측정", "주간 지표 기록, 저CTR 썸네일 교체", "KPI 시트 W2"),
        (5, "유통", "첫 앨범 10곡 제출, 팟캐스트 3편 등록, 스튜디오 RSS 연결", "발매 접수, 팟캐스트 탭"),
    ],
    3: [
        (0, "글로벌", "상위 10개 영상 영어·일본어 제목·설명 입력", "번역 메타 10편"),
        (1, "업로드", "롱폼 #3 공개 + 첫 7일 체크리스트 D-0 실행", "체크리스트 완료"),
        (2, "배포", "인스타 릴스·맘카페·블로그 임베드", "외부 유입 수"),
        (3, "커뮤니티", "Hype 요청 게시 (자격 충족 시) + 팁 카드", "하이프 수"),
        (4, "업로드", "롱폼 #4 공개 (해외 타깃, KST 09:00 예약)", "해외 조회 비율"),
        (5, "쇼츠", "쇼츠 3편 (영어 자막 포함)", "쇼츠 조회수"),
        (6, "측정", "초기 48시간 조회 추이 기록", "KPI 시트 W3"),
        (5, "유통", "YouTube Music 반영 확인, OAC 신청, Spotify·Apple 등록", "Topic 채널 URL"),
    ],
    4: [
        (0, "분석", "4주 지표 리뷰: 노출·CTR·AVD 전월 대비", "리뷰 노트"),
        (1, "분석", "A/B 승자 썸네일 패턴 확정 → 가이드 반영", "thumbnail_guide 갱신"),
        (2, "재작성", "저성과 영상 10편 제목·설명 재작성", "check_metadata.py 통과"),
        (3, "업로드", "롱폼 #5 공개 (승자 패턴 적용)", "CTR 48h"),
        (4, "커뮤니티", "후기 공유 + 다음 달 라인업 설문", "응답 수"),
        (5, "쇼츠", "쇼츠 2편", "쇼츠 조회수"),
        (6, "측정", "월간 KPI 마감, 다음 4주 계획 초안", "KPI 시트 W4"),
        (5, "유통", "트래픽 소스에 YouTube Music·팟캐스트 항목 생겼는지 확인", "KPI 시트 W4"),
    ],
}


def build_calendar(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "4주 캘린더"
    ws["A1"] = "아기 수면음악 채널 4주 실행 캘린더"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = "시작일(월요일)"
    ws["A2"].font = BOLD
    ws["B2"] = START
    ws["B2"].number_format = "yyyy-mm-dd"
    ws["B2"].fill = INPUT_FILL
    ws["B2"].font = INPUT_FONT
    ws["B2"].comment = Comment("시작 월요일을 바꾸면 모든 날짜가 자동 재계산됩니다.", "plan")
    ws["D2"] = "노란 배경·파란 글씨 셀만 수정하세요. 날짜는 B2에서 자동 계산됩니다."
    ws["D2"].font = Font(name=FONT, italic=True, color="666666")

    headers = ["주차", "주간 테마", "날짜", "요일", "유형", "작업", "산출물/측정", "담당", "완료(Y/N)", "메모"]
    ws.append([])
    ws.append(headers)
    hr = ws.max_row
    style_header(ws, hr, len(headers))

    days = ["월", "화", "수", "목", "금", "토", "일"]
    for week in range(1, 5):
        for wd, kind, task, output in sorted(PLAN[week], key=lambda t: t[0]):
            r = ws.max_row + 1
            offset = (week - 1) * 7 + wd
            ws.cell(r, 1, f"W{week}")
            ws.cell(r, 2, WEEK_THEMES[week])
            ws.cell(r, 3, f"=$B$2+{offset}")
            ws.cell(r, 3).number_format = "mm-dd"
            ws.cell(r, 4, days[wd])
            ws.cell(r, 5, kind)
            ws.cell(r, 6, task)
            ws.cell(r, 7, output)
            for c in (8, 9, 10):
                ws.cell(r, c).fill = INPUT_FILL
                ws.cell(r, c).font = INPUT_FONT
            for c in range(1, len(headers) + 1):
                cell = ws.cell(r, c)
                cell.border = BORDER
                if cell.font != INPUT_FONT:
                    cell.font = BODY_FONT
                cell.alignment = WRAP
    first_data = hr + 1
    last = ws.max_row
    # 예시 행: 담당·완료 형식 안내
    ws.cell(first_data, 8, "운영자")
    ws.cell(first_data, 9, "Y")
    ws.cell(first_data, 10, "예시 값입니다. 실제 진행에 맞춰 수정하세요.")

    # 요약
    ws.cell(last + 2, 1, "완료 건수").font = BOLD
    ws.cell(last + 2, 2, f'=COUNTIF(I{first_data}:I{last},"Y")')
    ws.cell(last + 3, 1, "전체 건수").font = BOLD
    ws.cell(last + 3, 2, f"=COUNTA(F{first_data}:F{last})")
    ws.cell(last + 4, 1, "진행률").font = BOLD
    ws.cell(last + 4, 2, f"=IF(B{last+3}=0,0,B{last+2}/B{last+3})")
    ws.cell(last + 4, 2).number_format = "0.0%"
    set_widths(ws, [6, 18, 8, 5, 9, 52, 26, 8, 10, 30])
    ws.freeze_panes = ws.cell(first_data, 1)

    # 주간 뷰
    wv = wb.create_sheet("주간 요약")
    wv.append(["주차", "테마", "업로드", "쇼츠", "커뮤니티", "측정 항목"])
    style_header(wv, 1, 6)
    for week in range(1, 5):
        r = week + 1
        wv.cell(r, 1, f"W{week}")
        wv.cell(r, 2, WEEK_THEMES[week])
        wv.cell(r, 3, f"=COUNTIFS('4주 캘린더'!$A${first_data}:$A${last},A{r},'4주 캘린더'!$E${first_data}:$E${last},\"업로드\")")
        wv.cell(r, 4, f"=COUNTIFS('4주 캘린더'!$A${first_data}:$A${last},A{r},'4주 캘린더'!$E${first_data}:$E${last},\"쇼츠\")")
        wv.cell(r, 5, f"=COUNTIFS('4주 캘린더'!$A${first_data}:$A${last},A{r},'4주 캘린더'!$E${first_data}:$E${last},\"커뮤니티\")")
        wv.cell(r, 6, ["기준선 노출·CTR·AVD", "CTR 48h, 쇼츠→롱폼 클릭", "해외 조회 비율, 하이프 수", "전월 대비 노출·AVD"][week - 1])
        for c in range(1, 7):
            wv.cell(r, c).border = BORDER
            wv.cell(r, c).font = BODY_FONT
    set_widths(wv, [6, 20, 8, 8, 10, 30])
    for ws_ in wb.worksheets:
        for row in ws_.iter_rows():
            for cell in row:
                if cell.font.name != FONT:
                    cell.font = Font(name=FONT, bold=cell.font.bold, color=cell.font.color, size=cell.font.size, italic=cell.font.italic)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # 열 때 전체 재계산
    wb.save(path)


# ---------------------------------------------------------------- KPI 추적
KPIS = [
    # (지표, 단위, 목표 규칙, 목표값, 설명)
    ("노출수", "회", "전월 대비 +50%", None, "스튜디오 > 도달범위 > 노출수. 목표는 기준선×1.5로 자동 계산"),
    ("노출 클릭률", "%", "4% 이상", 0.04, "스튜디오 > 도달범위. 소수로 입력 (4% = 0.04)"),
    ("평균 시청 지속 시간(분)", "분", "12분 이상", 12, "롱폼 기준. 스튜디오 > 참여도"),
    ("추천 동영상 트래픽 비율", "%", "40% 이상", 0.40, "트래픽 소스 중 '추천 동영상' 비율, 소수 입력"),
    ("재방문 시청자 비율", "%", "25% 이상", 0.25, "시청자층 > 재방문 시청자 / 전체"),
    ("구독자 순증", "명", "기준선 대비 증가", None, "시청자층 > 구독자"),
    ("쇼츠→롱폼 클릭", "회", "증가 추세", None, "쇼츠 관련 동영상 링크 클릭 (스튜디오 쇼츠 분석)"),
    ("라이브 평균 동시 시청자", "명", "증가 추세", None, "라이브 분석"),
    ("YouTube Music 트래픽 비율", "%", "생성 확인 후 증가 추세", None, "트래픽 소스 중 YouTube Music 비율, 소수 입력. 유통 반영 전에는 빈칸"),
    ("팟캐스트 트래픽 비율", "%", "생성 확인 후 증가 추세", None, "트래픽 소스 중 팟캐스트 비율, 소수 입력. RSS 연결 전에는 빈칸"),
]
WEEKS = ["W0 기준선", "W1", "W2", "W3", "W4"]


def build_kpi(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    ws["A1"] = "아기 수면음악 채널 KPI 추적 (28일 기준)"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = "노란 배경·파란 글씨 셀에 스튜디오 값을 입력하세요. 비율은 소수(4% = 0.04)로 입력합니다. 달성 여부·변화율은 자동 계산됩니다."
    ws["A2"].font = Font(name=FONT, italic=True, color="666666")
    headers = ["지표", "단위", "목표 규칙", "목표값"] + WEEKS + ["W4 vs W0", "달성", "설명"]
    ws.append([])
    ws.append(headers)
    hr = ws.max_row
    style_header(ws, hr, len(headers))
    first = hr + 1
    tgt_col = 4
    w0_col = 5
    w4_col = 9
    chg_col = 10
    ok_col = 11
    note_col = 12
    for i, (name, unit, rule, target, note) in enumerate(KPIS):
        r = first + i
        ws.cell(r, 1, name).font = BODY_FONT
        ws.cell(r, 2, unit).font = BODY_FONT
        ws.cell(r, 3, rule).font = BODY_FONT
        tcell = ws.cell(r, tgt_col)
        if target is None and name == "노출수":
            tcell.value = f"=IF(ISNUMBER({get_column_letter(w0_col)}{r}),{get_column_letter(w0_col)}{r}*1.5,\"\")"
            tcell.font = BODY_FONT
        elif target is None:
            tcell.value = f"=IF(ISNUMBER({get_column_letter(w0_col)}{r}),{get_column_letter(w0_col)}{r},\"\")"
            tcell.font = BODY_FONT
        else:
            tcell.value = target
            tcell.fill = INPUT_FILL
            tcell.font = INPUT_FONT
        for c in range(w0_col, w4_col + 1):
            cell = ws.cell(r, c)
            cell.fill = INPUT_FILL
            cell.font = INPUT_FONT
        fmt = "0.0%" if unit == "%" else "#,##0.0" if unit == "분" else "#,##0"
        for c in [tgt_col] + list(range(w0_col, w4_col + 1)):
            ws.cell(r, c).number_format = fmt
        w0 = f"{get_column_letter(w0_col)}{r}"
        w4 = f"{get_column_letter(w4_col)}{r}"
        tgt = f"{get_column_letter(tgt_col)}{r}"
        ws.cell(r, chg_col, f'=IF(AND(ISNUMBER({w0}),ISNUMBER({w4}),{w0}<>0),{w4}/{w0}-1,"")')
        ws.cell(r, chg_col).number_format = "+0.0%;-0.0%;0.0%"
        ws.cell(r, chg_col).font = BODY_FONT
        ws.cell(r, ok_col, f'=IF(OR(NOT(ISNUMBER({w4})),{tgt}=""),"",IF({w4}>={tgt},"달성","미달"))')
        ws.cell(r, ok_col).font = BODY_FONT
        ws.cell(r, note_col, note).font = Font(name=FONT, color="666666")
        for c in range(1, len(headers) + 1):
            ws.cell(r, c).border = BORDER
            ws.cell(r, c).alignment = WRAP
    last = first + len(KPIS) - 1
    # 예시 값: 노출수 행 W0 (현실적 규모 가정, 사용자 값으로 교체)
    ws.cell(first, w0_col, 120000)
    ws.cell(first, w0_col).comment = Comment("예시 값입니다. 스튜디오의 실제 28일 노출수로 교체하세요.", "plan")
    ws.cell(last + 2, 1, "달성 지표 수").font = BOLD
    ws.cell(last + 2, 2, f'=COUNTIF({get_column_letter(ok_col)}{first}:{get_column_letter(ok_col)}{last},"달성")')
    ws.cell(last + 3, 1, "평가 가능 지표 수").font = BOLD
    ws.cell(last + 3, 2, f'=COUNTIF({get_column_letter(ok_col)}{first}:{get_column_letter(ok_col)}{last},"달성")+COUNTIF({get_column_letter(ok_col)}{first}:{get_column_letter(ok_col)}{last},"미달")')
    set_widths(ws, [26, 6, 18, 12, 12, 10, 10, 10, 10, 11, 8, 48])
    ws.freeze_panes = ws.cell(first, 2)

    # 영상별 첫 7일 로그
    lg = wb.create_sheet("영상별 7일 로그")
    h2 = ["공개일", "제목", "시리즈", "썸네일 승자", "48h 노출", "48h CTR", "48h AVD(분)", "7일 조회", "7일 구독 전환", "하이프 수", "메모"]
    lg.append(h2)
    style_header(lg, 1, len(h2))
    lg.append([dt.date(2026, 9, 22), "아기 수면음악 | 잠 안 자는 신생아 5분 만에 재우는 자장가 10시간 광고없음", "밤잠 10시간", "A", 15000, 0.045, 14.2, 42000, 85, 12, "예시 행. 실제 값으로 교체"])
    for r in range(2, 32):
        for c in range(1, len(h2) + 1):
            cell = lg.cell(r, c)
            cell.fill = INPUT_FILL
            cell.font = INPUT_FONT
            cell.border = BORDER
        lg.cell(r, 1).number_format = "yyyy-mm-dd"
        lg.cell(r, 6).number_format = "0.0%"
        lg.cell(r, 7).number_format = "0.0"
    lg.cell(33, 4, "평균").font = BOLD
    for c, fmt in ((5, "#,##0"), (6, "0.0%"), (7, "0.0"), (8, "#,##0"), (9, "#,##0"), (10, "#,##0")):
        col = get_column_letter(c)
        lg.cell(33, c, f'=IF(COUNT({col}2:{col}31)=0,"",AVERAGE({col}2:{col}31))')
        lg.cell(33, c).number_format = fmt
        lg.cell(33, c).font = BOLD
    set_widths(lg, [11, 56, 12, 10, 10, 9, 11, 10, 12, 9, 30])
    lg.freeze_panes = "A2"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # 열 때 전체 재계산
    wb.save(path)


if __name__ == "__main__":
    (ROOT / "calendar").mkdir(exist_ok=True)
    build_calendar(ROOT / "calendar" / "4week_calendar.xlsx")
    build_kpi(ROOT / "tools" / "kpi_tracker.xlsx")
    print("built calendar/4week_calendar.xlsx and tools/kpi_tracker.xlsx")
